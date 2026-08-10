import re
from pathlib import Path
import openpyxl
from openpyxl.chart import PieChart, Reference

BASE_DIR = Path(__file__).resolve().parent

def add_pie_charts_to_sheet(ws):
    """
    Membuat 4 Pie Chart dengan tata letak Grid 2x2 agar tidak saling menumpuk.
    - Baris 1: Backend (F2)  | Web (M2)
    - Baris 2: Mobile (F17) | Total (M17)
    """
    categories = Reference(ws, min_col=2, min_row=6, max_col=4, max_row=6)

    # Konfigurasi penempatan grafik dengan jarak baris & kolom yang cukup
    charts_config = [
        {'title': 'Backend', 'row': 7,  'cell': 'F2'},
        {'title': 'Web',     'row': 8,  'cell': 'M2'},
        {'title': 'Mobile',  'row': 9,  'cell': 'F17'},
        {'title': 'Total',   'row': 10, 'cell': 'M17'}
    ]

    for cfg in charts_config:
        pie = PieChart()
        
        data = Reference(ws, min_col=2, min_row=cfg['row'], max_col=4, max_row=cfg['row'])
        
        pie.add_data(data, from_rows=True)
        pie.set_categories(categories)
        pie.title = cfg['title']
        
        # Dimensi grafik
        pie.width = 11.5
        pie.height = 6.2
        
        ws.add_chart(pie, cfg['cell'])


def sanitize_sheet_title(title: str) -> str:
    """Membersihkan karakter terlarang Excel dan membatasi max 31 karakter."""
    clean_title = re.sub(r'[\\/*?:\[\]]', '', str(title))
    return clean_title[:31]


def export_epic_report(template_name: str, output_name: str, epic_data: dict):
    template_path = BASE_DIR / "public" / template_name
    output_path = BASE_DIR / output_name

    wb = openpyxl.load_workbook(template_path)
    ws_template = wb['Template'] if 'Template' in wb.sheetnames else wb.active

    role_row_mapping = {
        'Backend': 7,
        'Web': 8,
        'Mobile': 9
    }

    stories = epic_data.get('stories', [])

    for story in stories:
        story_key = story.get('story_key', 'Sheet')
        clean_title = sanitize_sheet_title(story_key)
        
        ws = wb.copy_worksheet(ws_template)
        ws.title = clean_title

        # Header Metadata
        ws['B1'] = story.get('laporan_progres', '')
        ws['B2'] = story.get('epic_parent', '')
        ws['B3'] = story.get('status', '')

        # Populate Data per Role & Rumus Total per Baris (Kolom E)
        roles_data = story.get('roles', {})
        for role_name, row in role_row_mapping.items():
            counts = roles_data.get(role_name, {'to_do': 0, 'in_progress': 0, 'done': 0})
            ws[f'B{row}'] = counts.get('to_do', 0)
            ws[f'C{row}'] = counts.get('in_progress', 0)
            ws[f'D{row}'] = counts.get('done', 0)
            # Rumus Total Sub-task per Role (Kolom E)
            ws[f'E{row}'] = f"=SUM(B{row}:D{row})"

        # Rumus Total Keseluruhan per Status (Baris 10)
        ws['B10'] = "=SUM(B7:B9)"
        ws['C10'] = "=SUM(C7:C9)"
        ws['D10'] = "=SUM(D7:D9)"
        ws['E10'] = "=SUM(E7:E9)"

        # Generate Pie Charts
        add_pie_charts_to_sheet(ws)

    # Hapus sheet template bawaan
    if 'Template' in wb.sheetnames:
        wb.remove(wb['Template'])

    wb.save(output_path)
    print(f"File berhasil dibuat dengan {len(stories)} sheet di: {output_path}")


if __name__ == '__main__':
    sample_epic_payload = {
        'epic_key': 'JT-4',
        'stories': [
            {
                'story_key': 'JT-93',
                'laporan_progres': '[JT-93] Security Review - Enhance Modul Maintenance Taksonomi Data KUBL dan TKBI/KBLI',
                'epic_parent': '[JT-4] KBLI dan TKBI',
                'status': 'To Do',
                'roles': {
                    'Backend': {'to_do': 0, 'in_progress': 0, 'done': 0},
                    'Web':     {'to_do': 0, 'in_progress': 0, 'done': 0},
                    'Mobile':  {'to_do': 0, 'in_progress': 0, 'done': 0}
                }
            },
            {
                'story_key': 'JT-94',
                'laporan_progres': '[JT-94] Risk Management - Enhance Modul Maintenance Taksonomi Data KUBL dan TKBI/KBLI',
                'epic_parent': '[JT-4] KBLI dan TKBI',
                'status': 'In Progress',
                'roles': {
                    'Backend': {'to_do': 5, 'in_progress': 0, 'done': 0},
                    'Web':     {'to_do': 11, 'in_progress': 2, 'done': 2},
                    'Mobile':  {'to_do': 0, 'in_progress': 0, 'done': 0}
                }
            },
            {
                'story_key': 'JT-97',
                'laporan_progres': '[JT-97] Enhance Input Maintenance Pembiayaan Berkelanjutan Data KUBL',
                'epic_parent': '[JT-4] KBLI dan TKBI',
                'status': 'In Progress',
                'roles': {
                    'Backend': {'to_do': 2, 'in_progress': 3, 'done': 1},
                    'Web':     {'to_do': 4, 'in_progress': 1, 'done': 6},
                    'Mobile':  {'to_do': 1, 'in_progress': 0, 'done': 2}
                }
            }
        ]
    }

    export_epic_report(
        template_name='template_report.xlsx',
        output_name='Laporan_Progress_Epic_JT-4.xlsx',
        epic_data=sample_epic_payload
    )